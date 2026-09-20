import os
import openpyxl
import tkinter as tk
from tkinter import filedialog, messagebox

root = tk.Tk()
root.title("Excel Automation System")
root.geometry("500x320")

def select_file():
    file_path = filedialog.askopenfilename(
        title="Excel ဖိုင်ကို ရွေးချယ်ပါ",
        filetypes=[("Excel Files", "*.xlsx *.xls")]
    )
    if file_path:
        entry_file_path.delete(0, tk.END)
        entry_file_path.insert(0, file_path)

def start_process():
    path = entry_file_path.get()
    
    # 1. Validation Check
    if not path:
        messagebox.showwarning("သတိပေးချက်", "ကျေးဇူးပြု၍ Excel ဖိုင်ကို ပထမဦးစွာ ရွေးချယ်ပေးပါ!")
        return
        
    if not os.path.exists(path):
        messagebox.showerror("အမှားအယွင်း", "ဖိုင်လမ်းကြောင်း မမှန်ကန်ပါဗျာ!")
        return

    # 2. Excel Processing Logic
    try:
        # Excel Workbook ကို ဖွင့်ခြင်း
        wb = openpyxl.load_workbook(path)
        ws = wb.active

        # Header များ ထည့်သွင်းခြင်း
        ws['B1'] = "အသက်"
        ws['C1'] = "လုပ်သက်"

        # Formula များ ထည့်သွင်းခြင်း (DATEDIF Logic)
        ws['B2'] = '=DATEDIF(DATE(RIGHT(A2,4), MID(A2,4,2), LEFT(A2,2)), TODAY(), "Y") + 1'
        ws['C2'] = '=DATEDIF(DATE(RIGHT(A2,4), MID(A2,4,2), LEFT(A2,2)), TODAY(), "Y") & " နှစ်နှင့် " & DATEDIF(DATE(RIGHT(A2,4), MID(A2,4,2), LEFT(A2,2)), TODAY(), "YM") & " လ"'

        # ပြင်ဆင်ပြီးသော ဖိုင်ကို processed_report.xlsx အမည်ဖြင့် သိမ်းဆည်းခြင်း
        output_path = os.path.join(os.path.dirname(path), "processed_report.xlsx")
        wb.save(output_path)

        messagebox.showinfo(
            "အောင်မြင်ပါသည်။", 
            f"Excel Process ဆောင်ရွက်ပြီးပါပြီ!\n\nသိမ်းဆည်းခဲ့သည့်နေရာ -\n{output_path}"
        )

    except Exception as e:
        messagebox.showerror("Error", f"အောက်ပါ Error ဖြစ်ပွားခဲ့သည် -\n{str(e)}")

# --- UI Layout ရေးဆွဲခြင်း ---
label_title = tk.Label(root, text="ရုံးစာများ Excel သို့ ပြောင်းလဲပေးသည့် စနစ်", font=("Arial", 12, "bold"))
label_title.pack(pady=15)

frame_file = tk.Frame(root)
frame_file.pack(pady=10, padx=10, fill="x")

label_file = tk.Label(frame_file, text="Excel ဖိုင် -")
label_file.pack(side="left", padx=5)

entry_file_path = tk.Entry(frame_file, width=40)
entry_file_path.pack(side="left", padx=5)

btn_browse = tk.Button(frame_file, text="Browse...", command=select_file)
btn_browse.pack(side="left", padx=5)

btn_start = tk.Button(root, text="Process စတင်ရန်", command=start_process, bg="#4CAF50", fg="white", font=("Arial", 10, "bold"), height=2)
btn_start.pack(pady=20)

root.mainloop()